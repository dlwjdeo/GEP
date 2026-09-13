from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


DOC_PATH = Path(r"C:\Users\tlsal\OneDrive\바탕 화면\제작리포트.docx")


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "808080")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.first_child_found_in("w:shd")
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def main():
    doc = Document(DOC_PATH)
    target = None
    for table in doc.tables:
        if table.cell(0, 0).text.strip() == "획득 아이템":
            target = table
            break

    if target is None:
        raise RuntimeError("Could not find fishing probability table.")

    for row_idx, row in enumerate(target.rows):
        for cell in row.cells:
            set_cell_border(cell)
            if row_idx == 0:
                set_cell_shading(cell, "EDEDED")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Malgun Gothic"
                    run.font.size = Pt(9)
                    if row_idx == 0:
                        run.bold = True

    doc.save(DOC_PATH)


if __name__ == "__main__":
    main()
