from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


DOC_PATH = Path(r"C:\Users\tlsal\OneDrive\바탕 화면\제작리포트.docx")


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    inserted = Paragraph(new_p, paragraph._parent)
    if style:
        inserted.style = style
    if text:
        inserted.add_run(text)
    return inserted


def insert_table_after(paragraph, rows, widths):
    document = paragraph.part.document
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for row_idx, row in enumerate(rows):
        for col_idx, text in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = text
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Malgun Gothic"
                    run.font.size = Pt(9)
                    if row_idx == 0:
                        run.bold = True

    paragraph._p.addnext(table._tbl)
    return table


def set_paragraph_font(paragraph, size=10.5, bold=False):
    for run in paragraph.runs:
        run.font.name = "Malgun Gothic"
        run.font.size = Pt(size)
        run.bold = bold


def main():
    doc = Document(DOC_PATH)

    anchor = None
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "후보 위치에 가중치를 부여한 뒤 랜덤 선택하여 자연스럽게 배치":
            anchor = paragraph
            break

    if anchor is None:
        raise RuntimeError("Could not find item spawn section anchor paragraph.")

    inserted = []
    p = insert_paragraph_after(anchor, "낚시 자원 보정 확률")
    inserted.append((p, 10.5, True))

    lines = [
        "낚시 완료 시 물고기, 나무, 식물, 철 중 하나를 확률적으로 획득하도록 구성",
        "기본 확률은 물고기 50%, 나무 20%, 식물 20%, 철 10%로 시작",
        "철을 제외한 아이템이 나오면 해당 아이템 확률이 5% 감소하고 철 확률이 5% 증가",
        "철을 획득하면 확률이 기본값으로 초기화되어, 자원이 남아 있을 때 철을 노리는 클리어 루트가 가능",
    ]

    last = p
    for line in lines:
        last = insert_paragraph_after(last, line)
        inserted.append((last, 10.5, False))

    caption = insert_paragraph_after(last, "낚시 성공 시 기본 확률표")
    inserted.append((caption, 10.5, True))

    rows = [
        ["획득 아이템", "기본 확률", "획득 후 확률 변화"],
        ["물고기", "50%", "물고기 -5%, 철 +5%"],
        ["나무", "20%", "나무 -5%, 철 +5%"],
        ["식물", "20%", "식물 -5%, 철 +5%"],
        ["철", "10%", "전체 확률 기본값으로 초기화"],
    ]
    insert_table_after(caption, rows, [Inches(1.35), Inches(1.1), Inches(3.2)])

    for paragraph, size, bold in inserted:
        paragraph.paragraph_format.space_after = Pt(4)
        set_paragraph_font(paragraph, size=size, bold=bold)

    doc.save(DOC_PATH)


if __name__ == "__main__":
    main()
